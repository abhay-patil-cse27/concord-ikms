"""Document Parsers for PDF, DOCX, TXT, and Markdown files."""
import io
from abc import ABC, abstractmethod
from typing import List, Dict, Any
import pypdf
import docx


class ExtractedPage:
    def __init__(self, text: str, page_number: int, metadata: Dict[str, Any] = None):
        self.text = text
        self.page_number = page_number
        self.metadata = metadata or {}


class BaseDocumentParser(ABC):
    """Abstract contract for document content parsers."""

    @abstractmethod
    def extract_pages(self, file_bytes: bytes) -> List[ExtractedPage]:
        """Extract structured pages and text from raw file bytes."""
        pass


class PDFParser(BaseDocumentParser):
    """Extracts text from PDF documents page by page using pypdf, falling back to OCR on scanned pages."""

    def extract_pages(self, file_bytes: bytes) -> List[ExtractedPage]:
        from app.modules.documents.ocr import get_ocr_provider
        ocr_provider = get_ocr_provider()

        pages = []
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        total = len(reader.pages)

        for idx, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            is_ocr = False

            # If page text is very sparse (e.g. scanned image with no embedded font stream)
            if len(text) < 30:
                # Check for raster image objects on this page
                image_bytes_list = []
                try:
                    for img in getattr(page, "images", []):
                        if hasattr(img, "data") and img.data:
                            image_bytes_list.append(img.data)
                except Exception:
                    pass

                # If images were found or page text is completely empty, trigger OCR adapter
                if image_bytes_list or len(text) == 0:
                    ocr_result = ocr_provider.extract_text_from_images(image_bytes_list)
                    if ocr_result:
                        text = f"{text}\n\n{ocr_result}".strip() if text else ocr_result
                        is_ocr = True

            meta = {
                "total_pages": total,
                "is_ocr": is_ocr,
                "engine": "ocr" if is_ocr else "native_pdf"
            }
            pages.append(ExtractedPage(text=text.strip(), page_number=idx, metadata=meta))
        return pages


class DocxParser(BaseDocumentParser):
    """Extracts text and structured tables from Microsoft Word .docx documents."""

    def extract_pages(self, file_bytes: bytes) -> List[ExtractedPage]:
        doc = docx.Document(io.BytesIO(file_bytes))
        content_blocks = []
        table_count = len(doc.tables)

        for p in doc.paragraphs:
            if p.text.strip():
                content_blocks.append(p.text.strip())
                
        # Format Word tables as standard Markdown table matrices
        for t_idx, table in enumerate(doc.tables, start=1):
            if not table.rows:
                continue

            md_table_lines = []
            header_cells = [cell.text.strip() for cell in table.rows[0].cells]
            if header_cells and any(header_cells):
                md_table_lines.append("| " + " | ".join(header_cells) + " |")
                md_table_lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

            for row in table.rows[1:]:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):
                    md_table_lines.append("| " + " | ".join(row_cells) + " |")

            if md_table_lines:
                table_block = f"<!-- Table {t_idx} -->\n" + "\n".join(md_table_lines)
                content_blocks.append(table_block)

        combined_text = "\n\n".join(content_blocks)
        return [ExtractedPage(
            text=combined_text,
            page_number=1,
            metadata={
                "paragraph_count": len(doc.paragraphs),
                "table_count": table_count,
                "has_tables": table_count > 0
            }
        )]



class TextParser(BaseDocumentParser):
    """Extracts text from plain text or markdown files."""

    def extract_pages(self, file_bytes: bytes) -> List[ExtractedPage]:
        # Handle UTF-8 with fallback to latin-1
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="replace")
        return [ExtractedPage(text=text.strip(), page_number=1, metadata={})]


class ExcelParser(BaseDocumentParser):
    """Extracts structured rows, sheets, and cell data from Excel spreadsheets (.xlsx)."""

    def extract_pages(self, file_bytes: bytes) -> List[ExtractedPage]:
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        pages: List[ExtractedPage] = []

        for sheet_idx, sheet_name in enumerate(workbook.sheetnames, start=1):
            sheet = workbook[sheet_name]
            rows_text: List[str] = [f"=== Sheet: {sheet_name} ==="]

            for row in sheet.iter_rows(values_only=True):
                # Clean and filter non-empty cell values
                row_vals = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if row_vals:
                    rows_text.append(" | ".join(row_vals))

            if len(rows_text) > 1:  # More than just the header
                sheet_content = "\n".join(rows_text)
                pages.append(
                    ExtractedPage(
                        text=sheet_content,
                        page_number=sheet_idx,
                        metadata={"sheet_name": sheet_name, "row_count": len(rows_text) - 1}
                    )
                )

        if not pages:
            pages.append(ExtractedPage(text="[Empty Excel Workbook]", page_number=1, metadata={}))

        return pages


def get_parser_for_mime_or_ext(filename: str, mime_type: str) -> BaseDocumentParser:
    """Factory resolving appropriate parser based on file extension and MIME type."""
    fn_lower = filename.lower()
    if fn_lower.endswith(".pdf") or "pdf" in mime_type:
        return PDFParser()
    elif fn_lower.endswith(".docx") or "wordprocessingml" in mime_type:
        return DocxParser()
    elif fn_lower.endswith((".xlsx", ".xlsm", ".xltx")) or "spreadsheetml" in mime_type or "excel" in mime_type:
        return ExcelParser()
    elif fn_lower.endswith((".txt", ".md", ".csv", ".json", ".log")) or "text" in mime_type or "csv" in mime_type:
        return TextParser()
    else:
        # Default fallback to plain text parsing
        return TextParser()
