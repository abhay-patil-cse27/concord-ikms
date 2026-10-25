import io
import pytest
import docx
import pypdf
from app.modules.documents.parsers import PDFParser, DocxParser, TextParser, get_parser_for_mime_or_ext, ExtractedPage
from app.modules.documents.chunker import SlidingWindowTextChunker, TabularChunker, SemanticParagraphChunker, get_chunker_for_document


def test_text_parser():
    sample_text = "This is a test of the Institutional Knowledge Management System.\n\nSection 1: Guidelines."
    parser = TextParser()
    pages = parser.extract_pages(sample_text.encode("utf-8"))
    assert len(pages) == 1
    assert "Institutional Knowledge Management System" in pages[0].text
    assert pages[0].page_number == 1


def test_docx_parser():
    doc = docx.Document()
    doc.add_heading("Institutional Exam Policy", level=1)
    doc.add_paragraph("All semester examinations must be completed within 3 hours.")
    
    # Save to in-memory bytes
    f_bytes = io.BytesIO()
    doc.save(f_bytes)
    
    parser = DocxParser()
    pages = parser.extract_pages(f_bytes.getvalue())
    assert len(pages) == 1
    assert "Institutional Exam Policy" in pages[0].text
    assert "3 hours" in pages[0].text


def test_pdf_parser():
    # Create simple PDF in memory using pypdf
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    
    f_bytes = io.BytesIO()
    writer.write(f_bytes)
    
    parser = PDFParser()
    pages = parser.extract_pages(f_bytes.getvalue())
    assert len(pages) == 1
    assert pages[0].page_number == 1
    # Empty page triggers OCR fallback
    assert pages[0].metadata.get("is_ocr") is True
    assert "OCR extracted" in pages[0].text or len(pages[0].text) > 0


def test_ocr_provider_contract_and_mock():
    from app.modules.documents.ocr import BaseOCRProvider, set_ocr_provider, get_ocr_provider

    class MockTestOCR(BaseOCRProvider):
        def is_available(self) -> bool:
            return True
        def extract_text_from_images(self, image_bytes_list):
            return "MOCK_OCR_SCANNED_CONTENT_VERIFIED"

    original_provider = get_ocr_provider()
    try:
        set_ocr_provider(MockTestOCR())
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=200, height=200)
        f_bytes = io.BytesIO()
        writer.write(f_bytes)

        parser = PDFParser()
        pages = parser.extract_pages(f_bytes.getvalue())
        assert len(pages) == 1
        assert "MOCK_OCR_SCANNED_CONTENT_VERIFIED" in pages[0].text
        assert pages[0].metadata["is_ocr"] is True
    finally:
        set_ocr_provider(original_provider)


def test_chunker_sliding_window():
    from app.modules.documents.chunker import SlidingWindowTextChunker
    chunker = SlidingWindowTextChunker(chunk_size=100, chunk_overlap=20)
    long_text = (
        "Paragraph 1 contains some standard institutional background information that is quite extensive. "
        "Paragraph 2 continues with guidelines for attendance, assignments, and semester lab credits. "
        "Paragraph 3 concludes the policy document with regulatory references."
    )
    pages = [ExtractedPage(text=long_text, page_number=1)]
    chunks = chunker.chunk_pages(pages)
    
    assert len(chunks) > 1
    assert chunks[0].index == 0
    assert chunks[1].index == 1
    assert all(c.char_count <= 110 for c in chunks)
    assert all(c.metadata["page_number"] == 1 for c in chunks)


def test_tabular_chunker_preserves_table_structure():
    from app.modules.documents.chunker import TabularChunker
    table_text = (
        "=== Sheet: CSE_Faculty ===\n"
        "FacultyID | Name | Designation | Specialization\n"
        "F01 | Dr. R. Deshmukh | Professor | AI & NLP\n"
        "F02 | Prof. S. Patil | Asst. Professor | Cloud Systems\n"
        "F03 | Dr. M. Joshi | Assoc. Professor | Cyber Security\n"
    )
    pages = [ExtractedPage(text=table_text, page_number=1, metadata={"sheet_name": "CSE_Faculty"})]
    chunker = TabularChunker(rows_per_chunk=2)
    chunks = chunker.chunk_pages(pages)

    assert len(chunks) == 2
    assert "Columns: FacultyID" in chunks[0].content
    assert "Dr. R. Deshmukh" in chunks[0].content
    assert chunks[0].metadata["chunk_type"] == "tabular"
    assert chunks[1].metadata["row_start"] == 3


def test_chunker_factory():
    from app.modules.documents.chunker import (
        get_chunker_for_document,
        TabularChunker,
        SemanticParagraphChunker,
        SlidingWindowTextChunker
    )

    c_excel = get_chunker_for_document("grades.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    c_docx = get_chunker_for_document("curriculum.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    c_pdf = get_chunker_for_document("notice.pdf", "application/pdf")

    assert isinstance(c_excel, TabularChunker)
    assert isinstance(c_docx, SemanticParagraphChunker)
    assert isinstance(c_pdf, SlidingWindowTextChunker)


def test_parser_factory():
    p_pdf = get_parser_for_mime_or_ext("handbook.pdf", "application/pdf")
    p_docx = get_parser_for_mime_or_ext("curriculum.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    p_xlsx = get_parser_for_mime_or_ext("grades.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    p_txt = get_parser_for_mime_or_ext("notes.txt", "text/plain")

    from app.modules.documents.parsers import ExcelParser
    assert isinstance(p_pdf, PDFParser)
    assert isinstance(p_docx, DocxParser)
    assert isinstance(p_xlsx, ExcelParser)
    assert isinstance(p_txt, TextParser)


def test_excel_parser_extraction():
    import io
    import openpyxl
    from app.modules.documents.parsers import ExcelParser

    # Create an in-memory workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FeeStructure"
    ws.append(["Category", "Amount", "DueDate"])
    ws.append(["Tuition", 85000, "2026-08-30"])
    ws.append(["Library", 5000, "2026-08-30"])

    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    parser = ExcelParser()
    pages = parser.extract_pages(excel_bytes)

    assert len(pages) == 1
    assert "FeeStructure" in pages[0].text
    assert "Tuition | 85000 | 2026-08-30" in pages[0].text
    assert pages[0].metadata["row_count"] == 3


def test_docx_markdown_table_grid_extraction():
    """Verify DocxParser converts Word table matrices into clean Markdown tables."""
    doc = docx.Document()
    doc.add_heading("Autonomous Marks Distribution", level=1)
    doc.add_paragraph("The internal and end-semester mark breakdown is given below:")
    
    table = doc.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Max Marks"
    table.rows[0].cells[2].text = "Min Passing"

    table.rows[1].cells[0].text = "ISE (In-Semester)"
    table.rows[1].cells[1].text = "50"
    table.rows[1].cells[2].text = "20"

    table.rows[2].cells[0].text = "ESE (End-Semester)"
    table.rows[2].cells[1].text = "50"
    table.rows[2].cells[2].text = "20"

    f_bytes = io.BytesIO()
    doc.save(f_bytes)

    parser = DocxParser()
    pages = parser.extract_pages(f_bytes.getvalue())

    assert len(pages) == 1
    assert "| Component | Max Marks | Min Passing |" in pages[0].text
    assert "| --- | --- | --- |" in pages[0].text
    assert "| ESE (End-Semester) | 50 | 20 |" in pages[0].text
    assert pages[0].metadata.get("has_tables") is True


def test_semantic_paragraph_chunker_atomic_table_preservation():
    """Verify SemanticParagraphChunker preserves Markdown table blocks as atomic chunks."""
    markdown_doc = (
        "# Regulation 2024\n\n"
        "Here are the general passing guidelines for third-year engineering students.\n\n"
        "| Course Code | Course Name | Credits | Passing Threshold |\n"
        "| --- | --- | --- | --- |\n"
        "| CS301 | Database Engineering | 4 | 40% |\n"
        "| CS302 | Computer Networks | 4 | 40% |\n"
        "| CS303 | Artificial Intelligence | 4 | 40% |\n\n"
        "Students failing in more than 2 courses will be held back under KT regulations."
    )
    pages = [ExtractedPage(text=markdown_doc, page_number=1, metadata={"doc": "reg2024"})]
    chunker = SemanticParagraphChunker(max_chars=300)
    chunks = chunker.chunk_pages(pages)

    # Ensure table was kept as an atomic unit
    table_chunks = [c for c in chunks if c.metadata.get("has_table") is True]
    assert len(table_chunks) == 1
    assert "| CS301 | Database Engineering | 4 | 40% |" in table_chunks[0].content
    assert table_chunks[0].metadata["chunk_type"] == "table_grid"


def test_scanned_image_preprocessor_and_parser():
    """Verify ImagePreprocessor and ScannedImageParser on synthetic circular image."""
    from PIL import Image, ImageDraw
    from app.modules.documents.preprocessor import ImagePreprocessor
    from app.modules.documents.parsers import ScannedImageParser, get_parser_for_mime_or_ext

    # Create synthetic scanned college circular graphic (1000x400 white canvas with black text)
    img = Image.new("RGB", (1000, 400), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((50, 50), "KOLHAPUR INSTITUTE OF TECHNOLOGY", fill=(0, 0, 0))
    draw.text((50, 100), "CIRCULAR NO. KIT/ACAD/2026/044", fill=(0, 0, 0))
    draw.text((50, 150), "SUBJECT: SEMESTER EXAM TIME TABLE POSTPONED", fill=(0, 0, 0))

    img_bytes_io = io.BytesIO()
    img.save(img_bytes_io, format="PNG")
    raw_bytes = img_bytes_io.getvalue()

    # 1. Test ImagePreprocessor directly
    processed_img, meta = ImagePreprocessor.preprocess_image(raw_bytes)
    assert processed_img.mode == "1"  # Binarized 1-bit mode
    assert meta["is_preprocessed"] is True
    assert meta["original_dimensions"] == [1000, 400]
    assert meta["scale_factor"] == 1.5  # Upscaled for optimal OCR resolution

    # 2. Test Parser Resolution
    parser = get_parser_for_mime_or_ext("circular_scan_044.png", "image/png")
    assert isinstance(parser, ScannedImageParser)

    # 3. Test ScannedImageParser execution
    pages = parser.extract_pages(raw_bytes)
    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert pages[0].metadata.get("is_ocr") is True
    assert pages[0].metadata.get("media_type") == "scanned_image"
    assert "OCR" in pages[0].text or len(pages[0].text) > 0

